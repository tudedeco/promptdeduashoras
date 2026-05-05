from __future__ import annotations

import sqlite3
import threading
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook

from .grid import Cell

SCHEMA = """
CREATE TABLE IF NOT EXISTS cells (
  cell_id TEXT NOT NULL,
  query TEXT NOT NULL,
  parent_id TEXT,
  min_lat REAL, min_lon REAL, max_lat REAL, max_lon REAL,
  status TEXT CHECK(status IN ('pending','in_progress','done','subdivided','error')) NOT NULL DEFAULT 'pending',
  result_count INTEGER NOT NULL DEFAULT 0,
  hit_cap INTEGER NOT NULL DEFAULT 0,
  started_at TEXT,
  finished_at TEXT,
  error TEXT,
  PRIMARY KEY (query, cell_id)
);
CREATE INDEX IF NOT EXISTS idx_cells_q_status ON cells(query, status);

CREATE TABLE IF NOT EXISTS places (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  query TEXT NOT NULL,
  cid TEXT,
  dedupe_key TEXT NOT NULL,
  name TEXT,
  address TEXT,
  phone TEXT,
  rating REAL,
  reviews_count INTEGER,
  category TEXT,
  lat REAL,
  lon REAL,
  place_url TEXT,
  website TEXT,
  open_status TEXT,
  hours TEXT,
  card_text TEXT,
  cell_id TEXT,
  scraped_at TEXT,
  xlsx_exported INTEGER NOT NULL DEFAULT 0,
  UNIQUE(query, dedupe_key)
);
CREATE INDEX IF NOT EXISTS idx_places_query ON places(query);
CREATE INDEX IF NOT EXISTS idx_places_q_id ON places(query, id);
"""


@dataclass
class PlaceRecord:
    name: str | None
    address: str | None
    phone: str | None
    rating: float | None
    reviews_count: int | None
    category: str | None
    lat: float | None
    lon: float | None
    place_url: str | None
    cid: str | None
    website: str | None = None
    open_status: str | None = None
    hours: str | None = None
    card_text: str | None = None

    def dedupe_key(self) -> str:
        if self.cid:
            return f"cid:{self.cid}"
        lat_s = f"{self.lat:.4f}" if self.lat is not None else "?"
        lon_s = f"{self.lon:.4f}" if self.lon is not None else "?"
        name = (self.name or "").strip().lower()
        return f"nml:{name}|{lat_s}|{lon_s}"


XLSX_COLUMNS = [
    "name", "address", "phone", "rating", "reviews_count", "category",
    "lat", "lon", "place_url", "cid", "cell_id", "scraped_at", "query",
    "website", "open_status", "hours", "card_text",
]


class Store:
    """SQLite-backed state + streaming XLSX writer."""

    def __init__(self, db_path: Path, xlsx_path: Path, query: str):
        self.db_path = Path(db_path)
        self.xlsx_path = Path(xlsx_path)
        self.query = query
        self._lock = threading.RLock()       # protects SQLite
        self._xlsx_lock = threading.Lock()   # protects XLSX append
        self._conn = sqlite3.connect(self.db_path, check_same_thread=False, isolation_level=None)
        self._conn.execute("PRAGMA journal_mode=WAL")
        self._conn.execute("PRAGMA synchronous=NORMAL")
        self._conn.executescript(SCHEMA)
        self._ensure_schema_migrations()
        self._ensure_xlsx()
        try:
            self._flush_xlsx_pending()
        except Exception as e:
            print(f"[warn] xlsx flush failed at startup: {e}")

    def close(self) -> None:
        with self._lock:
            self._conn.close()

    def _ensure_xlsx(self) -> None:
        if self.xlsx_path.exists():
            with self._xlsx_lock:
                wb = load_workbook(self.xlsx_path)
                ws = wb["Places"] if "Places" in wb.sheetnames else wb.active
                existing = [cell.value for cell in ws[1]]
                missing = [column for column in XLSX_COLUMNS if column not in existing]
                for column in missing:
                    ws.cell(row=1, column=ws.max_column + 1, value=column)
                if missing:
                    wb.save(self.xlsx_path)
            return
        self._create_xlsx()

    def _create_xlsx(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Places"
        ws.append(XLSX_COLUMNS)
        wb.save(self.xlsx_path)

    def reset_query(self) -> tuple[int, int]:
        with self._lock:
            place_count = self._conn.execute(
                "SELECT COUNT(*) FROM places WHERE query=?",
                (self.query,),
            ).fetchone()[0]
            cell_count = self._conn.execute(
                "SELECT COUNT(*) FROM cells WHERE query=?",
                (self.query,),
            ).fetchone()[0]
            cur = self._conn.cursor()
            cur.execute("BEGIN")
            cur.execute("DELETE FROM places WHERE query=?", (self.query,))
            cur.execute("DELETE FROM cells WHERE query=?", (self.query,))
            cur.execute("COMMIT")
        with self._xlsx_lock:
            self._create_xlsx()
        return cell_count, place_count

    def _ensure_schema_migrations(self) -> None:
        cols = {row[1] for row in self._conn.execute("PRAGMA table_info(places)").fetchall()}
        if "xlsx_exported" not in cols:
            self._conn.execute("ALTER TABLE places ADD COLUMN xlsx_exported INTEGER NOT NULL DEFAULT 0")
        for col in ("website", "open_status", "hours", "card_text"):
            if col not in cols:
                self._conn.execute(f"ALTER TABLE places ADD COLUMN {col} TEXT")
        self._conn.execute("CREATE INDEX IF NOT EXISTS idx_places_q_xlsx ON places(query, xlsx_exported, id)")

    def seed_cells(self, cells: Iterable[Cell]) -> int:
        with self._lock:
            cur = self._conn.cursor()
            cur.execute("BEGIN")
            inserted = 0
            for c in cells:
                cur.execute(
                    """INSERT OR IGNORE INTO cells
                       (cell_id, query, parent_id, min_lat, min_lon, max_lat, max_lon, status)
                       VALUES (?,?,?,?,?,?,?, 'pending')""",
                    (c.cell_id, self.query, c.parent_id, c.min_lat, c.min_lon, c.max_lat, c.max_lon),
                )
                inserted += cur.rowcount
            cur.execute("COMMIT")
            return inserted

    def next_pending_cell(self) -> Cell | None:
        with self._lock:
            row = self._conn.execute(
                """UPDATE cells
                   SET status='in_progress', started_at=?
                   WHERE query=?
                     AND cell_id = (
                       SELECT cell_id
                       FROM cells
                       WHERE query=? AND status='pending'
                       ORDER BY max_lat DESC, min_lon ASC, (max_lat - min_lat) DESC, cell_id ASC
                       LIMIT 1
                     )
                   RETURNING cell_id, parent_id, min_lat, min_lon, max_lat, max_lon""",
                (_now(), self.query, self.query),
            ).fetchone()
            if not row:
                return None
            return Cell(cell_id=row[0], parent_id=row[1], min_lat=row[2], min_lon=row[3],
                        max_lat=row[4], max_lon=row[5])

    def mark_cell_done(self, cell_id: str, result_count: int, hit_cap: bool) -> None:
        with self._lock:
            self._conn.execute(
                """UPDATE cells SET status='done', result_count=?, hit_cap=?, finished_at=?
                   WHERE query=? AND cell_id=?""",
                (result_count, 1 if hit_cap else 0, _now(), self.query, cell_id),
            )

    def mark_cell_subdivided(self, cell_id: str, result_count: int, hit_cap: bool = True) -> None:
        with self._lock:
            self._conn.execute(
                """UPDATE cells SET status='subdivided', result_count=?, hit_cap=?, finished_at=?
                   WHERE query=? AND cell_id=?""",
                (result_count, 1 if hit_cap else 0, _now(), self.query, cell_id),
            )

    def mark_cell_error(self, cell_id: str, error: str) -> None:
        with self._lock:
            self._conn.execute(
                "UPDATE cells SET status='error', error=?, finished_at=? WHERE query=? AND cell_id=?",
                (error[:500], _now(), self.query, cell_id),
            )

    def reset_in_progress(self) -> int:
        with self._lock:
            cur = self._conn.execute(
                "UPDATE cells SET status='pending', started_at=NULL WHERE query=? AND status='in_progress'",
                (self.query,),
            )
            return cur.rowcount

    def cells_summary(self) -> dict[str, int]:
        with self._lock:
            rows = self._conn.execute(
                "SELECT status, COUNT(*) FROM cells WHERE query=? GROUP BY status",
                (self.query,),
            ).fetchall()
        return {status: count for status, count in rows}

    def count_places_in_cell(self, cell: Cell) -> int:
        with self._lock:
            return self._conn.execute(
                """SELECT COUNT(*)
                   FROM places
                   WHERE query=?
                     AND lat IS NOT NULL
                     AND lon IS NOT NULL
                     AND lat BETWEEN ? AND ?
                     AND lon BETWEEN ? AND ?""",
                (self.query, cell.min_lat, cell.max_lat, cell.min_lon, cell.max_lon),
            ).fetchone()[0]

    def upsert_places(self, places: list[PlaceRecord], cell_id: str) -> int:
        """Insert new places; return count of newly inserted rows."""
        inserted = 0
        if places:
            ts = _now()
            with self._lock:
                cur = self._conn.cursor()
                cur.execute("BEGIN")
                for p in places:
                    key = p.dedupe_key()
                    cur.execute(
                        """INSERT OR IGNORE INTO places
                           (query, cid, dedupe_key, name, address, phone, rating, reviews_count,
                            category, lat, lon, place_url, website, open_status, hours, card_text,
                            cell_id, scraped_at, xlsx_exported)
                           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,0)""",
                        (self.query, p.cid, key, p.name, p.address, p.phone, p.rating,
                         p.reviews_count, p.category, p.lat, p.lon, p.place_url, p.website,
                         p.open_status, p.hours, p.card_text, cell_id, ts),
                    )
                    inserted += cur.rowcount
                cur.execute("COMMIT")
        try:
            self._flush_xlsx_pending()
        except Exception as e:
            print(f"[warn] xlsx flush failed: {e}")
        return inserted

    def _append_xlsx(self, rows: list[tuple]) -> None:
        wb = load_workbook(self.xlsx_path)
        ws = wb["Places"]
        for r in rows:
            ws.append(r)
        wb.save(self.xlsx_path)

    def _flush_xlsx_pending(self, limit: int = 2000) -> int:
        total = 0
        with self._xlsx_lock:
            while True:
                with self._lock:
                    rows = self._conn.execute(
                        """SELECT id, name, address, phone, rating, reviews_count, category, lat, lon,
                                  place_url, cid, cell_id, scraped_at, query,
                                  website, open_status, hours, card_text
                           FROM places
                           WHERE query=? AND xlsx_exported=0
                           ORDER BY id ASC
                           LIMIT ?""",
                        (self.query, limit),
                    ).fetchall()
                if not rows:
                    return total
                row_ids = [row[0] for row in rows]
                xlsx_rows = [row[1:] for row in rows]
                self._append_xlsx(xlsx_rows)
                placeholders = ",".join("?" for _ in row_ids)
                with self._lock:
                    self._conn.execute(
                        f"UPDATE places SET xlsx_exported=1 WHERE query=? AND id IN ({placeholders})",
                        (self.query, *row_ids),
                    )
                total += len(row_ids)

    def cells_geojson(self) -> dict:
        with self._lock:
            rows = self._conn.execute(
                """SELECT cell_id, parent_id, min_lat, min_lon, max_lat, max_lon, status, result_count, hit_cap
                   FROM cells WHERE query=?""",
                (self.query,),
            ).fetchall()
        features = []
        for cid, pid, a, b, c, d, status, count, hit in rows:
            features.append({
                "type": "Feature",
                "properties": {
                    "cell_id": cid, "parent_id": pid, "status": status,
                    "result_count": count, "hit_cap": bool(hit),
                },
                "geometry": {
                    "type": "Polygon",
                    "coordinates": [[[b, a], [d, a], [d, c], [b, c], [b, a]]],
                },
            })
        return {"type": "FeatureCollection", "features": features}

    def places_since(self, last_id: int = 0, limit: int = 2000) -> tuple[list[dict], int]:
        with self._lock:
            rows = self._conn.execute(
                """SELECT id, name, address, phone, rating, reviews_count, category, lat, lon,
                          place_url, website, open_status
                   FROM places WHERE query=? AND id>? ORDER BY id ASC LIMIT ?""",
                (self.query, last_id, limit),
            ).fetchall()
        items = []
        max_id = last_id
        for rid, name, addr, phone, rating, reviews, cat, lat, lon, url, website, open_status in rows:
            items.append({
                "id": rid, "name": name, "address": addr, "phone": phone, "rating": rating,
                "reviews_count": reviews, "category": cat, "lat": lat, "lon": lon, "url": url,
                "website": website, "open_status": open_status,
            })
            if rid > max_id:
                max_id = rid
        return items, max_id


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")
